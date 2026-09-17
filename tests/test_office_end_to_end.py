"""End-to-end tests for Office formats conversion (OFF-001..003, OUT-001..003, CLI)."""

import subprocess
import sys
from pathlib import Path
import docx
import pypdfium2 as pdfium
import pytest

from openlargeprint.exporters import ExportOptions, PaperSize, PresetName
from openlargeprint.pipeline import PipelineOrchestrator
from test_docx_importer import create_sample_docx
from test_pptx_importer import create_sample_pptx


def test_docx_end_to_end_all_exporters(tmp_path: Path):
    """Verify DOCX can be converted to large-print DOCX, PDF, and Reader HTML (OFF-001, DOC-001)."""
    input_docx = tmp_path / "contract_source.docx"
    create_sample_docx(input_docx)

    orchestrator = PipelineOrchestrator()
    options = ExportOptions(preset=PresetName.LARGE, paper_size=PaperSize.A4)

    # 1. Convert to Large-Print DOCX
    out_docx = tmp_path / "output_large.docx"
    res_docx = orchestrator.convert(input_docx, out_docx, options=options, export_format="docx")
    assert res_docx.success
    assert out_docx.exists()
    doc = docx.Document(out_docx)
    assert len(doc.paragraphs) > 5
    assert len(doc.tables) >= 1

    # Verify font size is 20pt for standard paragraph
    body_runs = [r for p in doc.paragraphs for r in p.runs if r.font.size]
    assert any(r.font.size.pt == 20.0 for r in body_runs)

    # 2. Convert to Reflowed Large-Print PDF
    out_pdf = tmp_path / "output_large.pdf"
    res_pdf = orchestrator.convert(input_docx, out_pdf, options=options, export_format="pdf")
    assert res_pdf.success
    assert out_pdf.exists()
    pdf_doc = pdfium.PdfDocument(str(out_pdf))
    assert len(pdf_doc) >= 1
    # Verify A4 dimensions (595.28 x 841.89 pt)
    w, h = pdf_doc[0].get_size()
    assert abs(w - 595.28) < 1.0
    assert abs(h - 841.89) < 1.0

    # 3. Convert to In-App Reader HTML
    out_reader = tmp_path / "output_reader.html"
    res_reader = orchestrator.convert(input_docx, out_reader, options=options, export_format="reader")
    assert res_reader.success
    assert out_reader.exists()
    html_content = out_reader.read_text(encoding="utf-8")
    assert "<!DOCTYPE html>" in html_content
    assert "Contract Law Treatise" in html_content
    assert "large-print-table" in html_content


def test_pptx_end_to_end_all_exporters(tmp_path: Path):
    """Verify PPTX can be converted to large-print DOCX, PDF, and Reader HTML (OFF-001, DOC-001)."""
    input_pptx = tmp_path / "presentation_source.pptx"
    create_sample_pptx(input_pptx)

    orchestrator = PipelineOrchestrator()
    options = ExportOptions(preset=PresetName.COMFORTABLE, paper_size=PaperSize.A3)

    # 1. Convert to Large-Print DOCX
    out_docx = tmp_path / "output_slides.docx"
    res_docx = orchestrator.convert(input_pptx, out_docx, options=options, export_format="docx")
    assert res_docx.success
    assert out_docx.exists()
    doc = docx.Document(out_docx)
    assert any("Public International Law Lecture" in p.text for p in doc.paragraphs)
    assert any("[Speaker Notes]" in p.text for p in doc.paragraphs)

    # 2. Convert to Reflowed Large-Print PDF (A3)
    out_pdf = tmp_path / "output_slides_a3.pdf"
    res_pdf = orchestrator.convert(input_pptx, out_pdf, options=options, export_format="pdf")
    assert res_pdf.success
    assert out_pdf.exists()
    pdf_doc = pdfium.PdfDocument(str(out_pdf))
    assert len(pdf_doc) >= 1
    # Verify A3 dimensions (841.89 x 1190.55 pt)
    w, h = pdf_doc[0].get_size()
    assert abs(w - 841.89) < 1.0
    assert abs(h - 1190.55) < 1.0

    # 3. Convert to In-App Reader HTML
    out_reader = tmp_path / "output_slides.html"
    res_reader = orchestrator.convert(input_pptx, out_reader, options=options, export_format="reader")
    assert res_reader.success
    assert out_reader.exists()
    html_content = out_reader.read_text(encoding="utf-8")
    assert "Statute of the ICJ Article 38(1)" in html_content
    assert "Welcome students" in html_content


def test_office_cli_commands(tmp_path: Path):
    """Verify CLI convert and inspect commands work on DOCX and PPTX."""
    input_docx = tmp_path / "cli_sample.docx"
    create_sample_docx(input_docx)

    out_pdf = tmp_path / "cli_output.pdf"

    # Test CLI convert
    cmd_convert = [
        sys.executable,
        "-m",
        "openlargeprint.cli",
        "convert",
        str(input_docx),
        "-o",
        str(out_pdf),
        "--preset",
        "Large",
    ]
    res_c = subprocess.run(cmd_convert, capture_output=True, text=True)
    assert res_c.returncode == 0
    assert out_pdf.exists()
    assert "Successfully converted" in res_c.stdout

    # Test CLI inspect
    cmd_inspect = [
        sys.executable,
        "-m",
        "openlargeprint.cli",
        "inspect",
        str(input_docx),
        "--json",
    ]
    res_i = subprocess.run(cmd_inspect, capture_output=True, text=True)
    assert res_i.returncode == 0
    assert '"schema_version": "1.0.0"' in res_i.stdout
