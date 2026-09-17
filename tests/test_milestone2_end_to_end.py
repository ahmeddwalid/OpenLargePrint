"""End-to-end integration test proving Milestone 2:
Scanned PDF -> PaddleOCR layout & recognition -> DocumentIR -> UNMODIFIED DocxExporter.
"""

from pathlib import Path
from PIL import Image, ImageDraw
import docx
from docx.shared import Pt
import pytest
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from openlargeprint.exporters import ExportOptions, PaperSize, PresetName
from openlargeprint.ir.models import BlockType, ExtractionMethod, PageClassification
from openlargeprint.pipeline import PipelineOrchestrator


def generate_scanned_law_page(pdf_path: Path, img_path: Path):
    """Generate a realistic scanned page image and wrap it in a PDF with zero text layer."""
    w_px, h_px = 1700, 2200
    img = Image.new("RGB", (w_px, h_px), color="white")
    draw = ImageDraw.Draw(img)

    # Header across top
    draw.text((150, 120), "LAW OF TORTS: VICARIOUS LIABILITY", fill="black")

    # Column 1
    draw.text((150, 250), "1. The Employment Relationship", fill="black")
    draw.text((150, 330), "An employer is strictly liable for torts committed in the course", fill="black")
    draw.text((150, 390), "of employment, regardless of personal fault or negligence.", fill="black")
    draw.text((150, 470), "The modern test considers integration, control, and enterprise risk.", fill="black")
    draw.text((150, 550), "- Close connection doctrine", fill="black")
    draw.text((150, 610), "- Salmond formulation scope", fill="black")

    # Column 2
    draw.text((950, 250), "2. Independent Contractors", fill="black")
    draw.text((950, 330), "General rule: no liability for torts of independent contractors.", fill="black")
    draw.text((950, 390), "Exceptions arise in non-delegable duty or extra-hazardous acts.", fill="black")
    draw.text((950, 470), "Woodland v Essex County Council clarifies educational duty scope.", fill="black")

    img.save(img_path)

    # Wrap in PDF
    c = canvas.Canvas(str(pdf_path), pagesize=letter)
    c.drawImage(str(img_path), 0, 0, width=letter[0], height=letter[1])
    c.showPage()
    c.save()


def test_milestone2_scanned_pdf_to_docx_end_to_end(tmp_path: Path):
    """Execute complete Milestone 2 vertical slice:
    Scanned PDF -> OCR recognition -> DocumentIR -> unmodified DocxExporter -> 20pt DOCX.
    """
    pdf_file = tmp_path / "torts_scanned.pdf"
    img_file = tmp_path / "torts_scanned.png"
    out_docx = tmp_path / "torts_large_print.docx"

    generate_scanned_law_page(pdf_file, img_file)

    orchestrator = PipelineOrchestrator()
    options = ExportOptions(
        preset=PresetName.LARGE,  # 20pt body, 1.5 line spacing (OUT-006)
        paper_size=PaperSize.A4,   # A4 paper size (OUT-007)
        include_page_markers=True, # Original page markers (OUT-005)
    )

    result = orchestrator.convert(pdf_file, out_docx, options)

    # 1. Pipeline result assertions
    assert result.success is True
    assert out_docx.exists()
    assert out_docx.stat().st_size > 0

    doc_ir = result.document_ir

    # 2. Canonical DocumentIR assertions under OCR (DOC-001, PDF-003)
    assert doc_ir.schema_version == "1.0.0"
    assert doc_ir.metadata.page_count == 1
    assert doc_ir.pages[0].classification == PageClassification.SCANNED

    ocr_blocks = [b for b in doc_ir.blocks if b.type != BlockType.PAGE_MARKER]
    assert len(ocr_blocks) > 0

    # Every OCR block must carry OCR_FAST extraction method and bounding box in PDF points (PDF-006)
    for block in ocr_blocks:
        assert block.extraction_method == ExtractionMethod.OCR_FAST
        assert block.confidence > 0.60
        assert block.source_bounding_box is not None
        assert block.source_page == 1

    # Verify column reading order preserved under OCR (PDF-003)
    texts = [b.text for b in ocr_blocks if b.text]
    full_text = " ".join(texts)
    assert "VICARIOUS LIABILITY" in full_text or "LAW OF TORTS" in full_text

    emp_pos = next(i for i, t in enumerate(texts) if "Employ" in t or "employer" in t)
    indep_pos = next(i for i, t in enumerate(texts) if "Independent Contractors" in t or "contractors" in t)
    assert emp_pos < indep_pos, "Column 1 text must precede Column 2 text in OCR reading order"

    # 3. Unmodified DOCX Exporter assertions (OUT-001, OUT-006, OUT-007, OUT-009)
    doc = docx.Document(str(out_docx))
    section = doc.sections[0]

    # Verify A4 physical dimensions embedded (OUT-007, OUT-009)
    assert pytest.approx(section.page_width.mm, 1.0) == 210.0
    assert pytest.approx(section.page_height.mm, 1.0) == 297.0

    # Verify print reminder in metadata
    assert "Print at 100%" in doc.core_properties.comments

    # Verify 20pt body font and 1.5 line spacing applied to recognized OCR text
    body_paras = [p for p in doc.paragraphs if "liable" in p.text or "employer" in p.text]
    assert len(body_paras) > 0
    first_body = body_paras[0]
    assert first_body.runs[0].font.size == Pt(20)
    assert first_body.paragraph_format.line_spacing == 1.5

    # Verify single-column reflow
    assert len(doc.tables) == 0
