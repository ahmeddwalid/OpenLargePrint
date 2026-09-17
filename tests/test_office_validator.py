"""Tests for content-based sniffing of Office formats (DOCX, PPTX, DOC, PPT) (SEC-001)."""

import zipfile
from pathlib import Path
import docx
import pptx
import pytest

from openlargeprint.security import (
    SecurityValidationError,
    detect_file_type,
)
from openlargeprint.security.validator import (
    OLE2_MAGIC,
    OLE2_PPT_STREAM_1,
    OLE2_WORD_STREAM,
)


def test_detect_docx_by_content(tmp_path: Path):
    """Verify DOCX is recognized by OpenXML content regardless of file extension (SEC-001)."""
    doc_path = tmp_path / "sample_doc.bin"
    doc = docx.Document()
    doc.add_paragraph("This is a Word document.")
    doc.save(str(doc_path))

    detected = detect_file_type(doc_path)
    assert detected == "docx"


def test_detect_pptx_by_content(tmp_path: Path):
    """Verify PPTX is recognized by OpenXML content regardless of file extension (SEC-001)."""
    ppt_path = tmp_path / "presentation.dat"
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Presentation Title"
    prs.save(str(ppt_path))

    detected = detect_file_type(ppt_path)
    assert detected == "pptx"


def test_detect_legacy_doc_by_content(tmp_path: Path):
    """Verify legacy binary .doc is recognized by OLE2 magic bytes and Word stream (SEC-001)."""
    doc_file = tmp_path / "legacy.doc"
    # Construct minimal synthetic OLE2 container containing WordDocument stream
    data = bytearray(OLE2_MAGIC)
    data.extend(b"\x00" * 512)
    data.extend(OLE2_WORD_STREAM)
    data.extend(b"\x00" * 1024)
    doc_file.write_bytes(data)

    detected = detect_file_type(doc_file)
    assert detected == "doc"


def test_detect_legacy_ppt_by_content(tmp_path: Path):
    """Verify legacy binary .ppt is recognized by OLE2 magic bytes and PowerPoint stream (SEC-001)."""
    ppt_file = tmp_path / "legacy.ppt"
    # Construct minimal synthetic OLE2 container containing PowerPoint Document stream
    data = bytearray(OLE2_MAGIC)
    data.extend(b"\x00" * 512)
    data.extend(OLE2_PPT_STREAM_1)
    data.extend(b"\x00" * 1024)
    ppt_file.write_bytes(data)

    detected = detect_file_type(ppt_file)
    assert detected == "ppt"


def test_reject_arbitrary_zip_archive(tmp_path: Path):
    """Verify an arbitrary ZIP archive is rejected if it lacks Word/PowerPoint structures (SEC-001)."""
    zip_path = tmp_path / "random_archive.zip"
    with zipfile.ZipFile(zip_path, "w") as zf:
        zf.writestr("test.txt", "Some arbitrary file")

    with pytest.raises(SecurityValidationError, match="does not contain a valid Word"):
        detect_file_type(zip_path)


def test_reject_arbitrary_ole2_file(tmp_path: Path):
    """Verify an arbitrary OLE2 file is rejected if it lacks Word/PPT streams."""
    ole_path = tmp_path / "excel_or_other.ole"
    data = bytearray(OLE2_MAGIC)
    data.extend(b"\x00" * 512)
    data.extend("Workbook".encode("utf-16le"))
    data.extend(b"\x00" * 512)
    ole_path.write_bytes(data)

    with pytest.raises(SecurityValidationError, match="not a supported Word"):
        detect_file_type(ole_path)
