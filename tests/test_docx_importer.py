"""Unit tests for DocxImporter (OFF-001, DOC-001, LANG-001, SEC-001)."""

import io
from pathlib import Path
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image
import pytest

from openlargeprint.importers.office.docx import DocxImporter
from openlargeprint.ir.models import BlockType, TextDirection
from openlargeprint.ir.validator import validate_document_ir
from openlargeprint.security import JobWorkspace


def create_sample_docx(file_path: Path) -> Path:
    """Create a rich DOCX document with headings, lists, tables, images, and Arabic RTL."""
    doc = docx.Document()

    # 1. Document Title
    doc.add_heading("Contract Law Treatise: Principles and Remedies", level=0)

    # 2. Heading 1
    doc.add_heading("1. Formation of Valid Contracts", level=1)

    # 3. Standard Paragraph
    doc.add_paragraph("A contract requires offer, acceptance, consideration, and intention to create legal relations.")

    # 4. Heading 2
    doc.add_heading("1.1 Consideration Doctrine", level=2)

    # 5. List items
    doc.add_paragraph("Past consideration is not good consideration (Roscorla v Thomas).", style="List Bullet")
    doc.add_paragraph("Performance of an existing duty is insufficient unless practical benefit exists.", style="List Bullet")

    # 6. Quote
    p_quote = doc.add_paragraph("Consideration must be sufficient, but need not be adequate.", style="Quote")

    # 7. Embedded Table
    tbl = doc.add_table(rows=3, cols=3)
    hdr_cells = tbl.rows[0].cells
    hdr_cells[0].text = "Doctrine"
    hdr_cells[1].text = "Leading Case"
    hdr_cells[2].text = "Year"

    row1 = tbl.rows[1].cells
    row1[0].text = "Promissory Estoppel"
    row1[1].text = "High Trees House"
    row1[2].text = "1947"

    row2 = tbl.rows[2].cells
    row2[0].text = "Practical Benefit"
    row2[1].text = "Williams v Roffey Bros"
    row2[2].text = "1991"

    # 8. Caption
    doc.add_paragraph("Table 1.1: Landmark Consideration Authorities", style="Caption")

    # 9. Embedded Image
    img_buf = io.BytesIO()
    img = Image.new("RGB", (200, 100), color=(70, 130, 180))
    img.save(img_buf, format="PNG")
    img_buf.seek(0)
    doc.add_picture(img_buf, width=Inches(2.0))

    # 10. Arabic RTL Paragraph (LANG-001)
    p_ar = doc.add_paragraph("العقد شريعة المتعاقدين فلا يجوز نقضه ولا تعديله إلا باتفاق الطرفين.")
    pPr = p_ar._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:bidi"))

    # 11. Hard Page Break
    doc.add_page_break()

    # 12. Page 2 Content
    doc.add_heading("2. Vitiating Factors", level=1)
    doc.add_paragraph("A contract may be voidable due to misrepresentation, duress, or undue influence.")

    # 13. Footnote style paragraph
    try:
        doc.styles.add_style("Footnote Text", docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
    except Exception:
        pass
    doc.add_paragraph("1 Chappell & Co Ltd v Nestle Co Ltd [1960] AC 87.", style="Footnote Text")

    doc.save(str(file_path))
    return file_path


def test_docx_importer_structure(tmp_path: Path):
    """Verify DOCX document is imported with structure, hierarchy, and RTL preserved (OFF-001, DOC-001)."""
    docx_file = tmp_path / "sample_contract.docx"
    create_sample_docx(docx_file)

    importer = DocxImporter()
    with JobWorkspace() as ws:
        doc_ir = importer.import_document(docx_file, ws)

        # Check image extraction while workspace is alive
        images = [b for b in doc_ir.blocks if b.type == BlockType.IMAGE]
        assert len(images) == 1
        assert images[0].image_asset is not None
        assert images[0].image_asset.width == 200
        assert images[0].image_asset.height == 100
        assert Path(images[0].image_asset.file_path).exists()

    # Validate canonical schema
    warnings = validate_document_ir(doc_ir)
    assert not warnings, f"DocumentIR validation failed: {warnings}"

    # Verify pages
    assert len(doc_ir.pages) == 2
    assert doc_ir.metadata.page_count == 2
    assert "Contract Law Treatise" in doc_ir.metadata.title

    # Verify block types
    types = [b.type for b in doc_ir.blocks]
    assert BlockType.PAGE_MARKER in types
    assert BlockType.HEADING in types
    assert BlockType.PARAGRAPH in types
    assert BlockType.LIST in types
    assert BlockType.QUOTE in types
    assert BlockType.TABLE in types
    assert BlockType.CAPTION in types
    assert BlockType.IMAGE in types
    assert BlockType.FOOTNOTE in types

    # Check headings
    headings = [b for b in doc_ir.blocks if b.type == BlockType.HEADING]
    assert len(headings) >= 3
    assert headings[0].level in (1, 2)
    assert any("Formation of Valid Contracts" in h.text for h in headings)

    # Check table structure
    tables = [b for b in doc_ir.blocks if b.type == BlockType.TABLE]
    assert len(tables) == 1
    t = tables[0].table_structure
    assert t is not None
    assert t.row_count == 3
    assert t.column_count == 3
    assert t.rows[0][0].text == "Doctrine"
    assert t.rows[1][1].text == "High Trees House"

    # Check Arabic RTL block (LANG-001)
    ar_blocks = [b for b in doc_ir.blocks if b.language == "ar"]
    assert len(ar_blocks) >= 1
    assert ar_blocks[0].text_direction == TextDirection.RTL
    assert "العقد شريعة المتعاقدين" in ar_blocks[0].text


def test_docx_importer_invalid_file(tmp_path: Path):
    """Verify corrupted or invalid DOCX raises descriptive ValueError."""
    bad_file = tmp_path / "corrupted.docx"
    bad_file.write_bytes(b"PK\x03\x04not a real zip")

    importer = DocxImporter()
    with JobWorkspace() as ws:
        with pytest.raises(ValueError, match="Failed to parse DOCX document"):
            importer.import_document(bad_file, ws)
