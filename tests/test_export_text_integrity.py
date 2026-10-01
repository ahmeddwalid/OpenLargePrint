"""Export typography must not rewrite source characters (DOC-002, OUT-003)."""
from pathlib import Path
import docx
import pypdfium2 as pdfium
from openlargeprint.exporters import PdfExporter, DocxExporter, ReaderExporter, ExportOptions
from openlargeprint.ir.models import Block, BlockType, DocumentIR, DocumentMetadata
from openlargeprint.ir.models import TableCell, TableStructure
from openlargeprint.ir.models import TextDirection


def test_full_height_retained_table_exports_at_both_sizes(tmp_path):
    from PIL import Image
    from openlargeprint.ir.models import ImageAsset
    from openlargeprint.exporters import PresetName
    source = tmp_path / "table.png"
    Image.new("RGB", (1000, 1490), "white").save(source)
    table = TableStructure(rows=[[TableCell(text="Heading")], [TableCell(text="Value")]], has_header=True)
    block = Block(id="table", type=BlockType.TABLE, source_page=294,
        table_structure=table, warnings=["Compare with the original table."],
        image_asset=ImageAsset(asset_id="table_source", file_path=str(source),
            width=1000, height=1490, mime_type="image/png"))
    ir = DocumentIR(metadata=DocumentMetadata(title="Retained table"), blocks=[block])
    for preset in (PresetName.LARGE, PresetName.VERY_LARGE):
        path = PdfExporter().export(ir, tmp_path / f"{preset.value}.pdf", ExportOptions(preset=preset))
        with pdfium.PdfDocument(path) as pdf:
            assert len(pdf) >= 1


def test_list_markers_survive_every_export(tmp_path: Path):
    text = "* A B C well- known."
    ir = DocumentIR(metadata=DocumentMetadata(title="Preservation"), blocks=[
        Block(id="list", type=BlockType.LIST, text=text, source_page=1)])
    pdf_path = PdfExporter().export(ir, tmp_path / "list.pdf", ExportOptions())
    with pdfium.PdfDocument(pdf_path) as pdf:
        page = pdf[0].get_textpage()
        try:
            assert text in page.get_text_range()
        finally:
            page.close()
    docx_path = DocxExporter().export(ir, tmp_path / "list.docx", ExportOptions())
    assert text in "\n".join(p.text for p in docx.Document(docx_path).paragraphs)
    html_path = ReaderExporter().export(ir, tmp_path / "list.html", ExportOptions())
    assert text in html_path.read_text(encoding="utf-8")


def test_merged_cells_survive_every_export(tmp_path):
    table = TableStructure(rows=[
        [TableCell(text="Shared heading", col_span=2), TableCell()],
        [TableCell(text="Left"), TableCell(text="Right")]], has_header=True)
    block = Block(id="table", type=BlockType.TABLE, table_structure=table, source_page=1)
    ir = DocumentIR(metadata=DocumentMetadata(title="Merged cells"), blocks=[block])
    html_path = ReaderExporter().export(ir, tmp_path / "merged.html")
    assert 'colspan="2"' in html_path.read_text(encoding="utf-8")
    path = DocxExporter().export(ir, tmp_path / "merged.docx")
    assert 'w:gridSpan w:val="2"' in docx.Document(path).tables[0]._tbl.xml
    exporter = PdfExporter()
    flow = exporter._build_platypus_table(table, ExportOptions(), exporter._create_typography_styles(ExportOptions()), 470, False)
    assert ('SPAN', (0, 0), (1, 0)) in flow._spanCmds


def test_rtl_merged_anchor_is_not_dropped(tmp_path):
    table = TableStructure(rows=[
        [TableCell(text="Shared heading", col_span=2), TableCell()],
        [TableCell(text="Left"), TableCell(text="Right")]], has_header=True)
    ir = DocumentIR(metadata=DocumentMetadata(title="RTL merged cells"), blocks=[
        Block(id="table", type=BlockType.TABLE, table_structure=table,
              source_page=1, text_direction=TextDirection.RTL)])
    path = PdfExporter().export(ir, tmp_path / "rtl.pdf")
    with pdfium.PdfDocument(path) as pdf:
        page = pdf[0].get_textpage()
        try:
            assert "Shared heading" in page.get_text_range()
        finally:
            page.close()
