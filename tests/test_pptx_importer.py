"""Unit tests for PptxImporter (OFF-001, DOC-001, LANG-001, SEC-001)."""

import io
from pathlib import Path
from PIL import Image
import pptx
from pptx.util import Inches, Pt
import pytest

from openlargeprint.importers.office.pptx import PptxImporter
from openlargeprint.ir.models import BlockType, TextDirection
from openlargeprint.ir.validator import validate_document_ir
from openlargeprint.security import JobWorkspace


def create_sample_pptx(file_path: Path) -> Path:
    """Create a multi-slide PPTX presentation with titles, bullets, tables, images, and notes."""
    prs = pptx.Presentation()

    # --- Slide 1: Title Slide ---
    slide_layout_0 = prs.slide_layouts[0]
    slide1 = prs.slides.add_slide(slide_layout_0)
    slide1.shapes.title.text = "Public International Law Lecture"
    slide1.placeholders[1].text = "Sources of International Law and State Responsibility"

    # Slide 1 Speaker Notes (OFF-001)
    notes_slide1 = slide1.notes_slide
    notes_slide1.notes_text_frame.text = "Welcome students and distribute syllabus."

    # --- Slide 2: Content, Table, Image, Bullets ---
    slide_layout_blank = prs.slide_layouts[6]
    slide2 = prs.slides.add_slide(slide_layout_blank)

    # Manual Title Shape
    title_box = slide2.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    title_box.text_frame.text = "Statute of the ICJ Article 38(1)"

    # Bulleted List Shape
    list_box = slide2.shapes.add_textbox(Inches(1), Inches(1.8), Inches(4), Inches(2))
    tf = list_box.text_frame
    p1 = tf.paragraphs[0]
    p1.text = "International conventions and treaties"
    p1.level = 0
    p2 = tf.add_paragraph()
    p2.text = "International custom as evidence of general practice"
    p2.level = 1

    # Table Shape
    table_shape = slide2.shapes.add_table(rows=2, cols=2, left=Inches(1), top=Inches(4.2), width=Inches(4), height=Inches(1.2))
    tbl = table_shape.table
    tbl.rows[0].cells[0].text = "Source"
    tbl.rows[0].cells[1].text = "Authority"
    tbl.rows[1].cells[0].text = "Treaty"
    tbl.rows[1].cells[1].text = "UN Charter Art 2(4)"

    # Image Shape
    img_buf = io.BytesIO()
    Image.new("RGB", (160, 80), color=(34, 139, 34)).save(img_buf, format="PNG")
    img_buf.seek(0)
    slide2.shapes.add_picture(img_buf, left=Inches(5.5), top=Inches(1.8), width=Inches(2.5))

    # Slide 2 Speaker Notes
    notes_slide2 = slide2.notes_slide
    notes_slide2.notes_text_frame.text = "Emphasize distinction between hard law and soft law."

    # --- Slide 3: Arabic RTL Slide (LANG-001) ---
    slide3 = prs.slides.add_slide(slide_layout_blank)
    ar_box = slide3.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    ar_tf = ar_box.text_frame
    ar_tf.text = "مصادر القانون الدولي العام وفقاً لمحكمة العدل الدولية"

    prs.save(str(file_path))
    return file_path


def test_pptx_importer_structure(tmp_path: Path):
    """Verify PPTX document is imported with slides, hierarchy, tables, images, and notes (OFF-001)."""
    pptx_file = tmp_path / "sample_lecture.pptx"
    create_sample_pptx(pptx_file)

    importer = PptxImporter()
    with JobWorkspace() as ws:
        doc_ir = importer.import_document(pptx_file, ws)

        # Check image extraction while workspace is alive
        images = [b for b in doc_ir.blocks if b.type == BlockType.IMAGE]
        assert len(images) == 1
        assert images[0].image_asset is not None
        assert images[0].image_asset.width == 160
        assert images[0].image_asset.height == 80
        assert Path(images[0].image_asset.file_path).exists()

    # Validate canonical schema
    warnings = validate_document_ir(doc_ir)
    assert not warnings, f"DocumentIR validation failed: {warnings}"

    # Verify pages/slides count
    assert len(doc_ir.pages) == 3
    assert doc_ir.metadata.page_count == 3
    assert "Public International Law Lecture" in doc_ir.metadata.title

    # Verify slide markers
    page_markers = [b for b in doc_ir.blocks if b.type == BlockType.PAGE_MARKER]
    assert len(page_markers) == 3
    assert [m.page_marker for m in page_markers] == [1, 2, 3]

    # Verify headings
    headings = [b for b in doc_ir.blocks if b.type == BlockType.HEADING]
    assert len(headings) >= 1
    assert "Public International Law Lecture" in [h.text for h in headings]

    # Verify bulleted list items
    lists = [b for b in doc_ir.blocks if b.type == BlockType.LIST]
    assert len(lists) >= 1
    assert any("International custom" in l.text for l in lists)

    # Verify table structure
    tables = [b for b in doc_ir.blocks if b.type == BlockType.TABLE]
    assert len(tables) == 1
    t = tables[0].table_structure
    assert t is not None
    assert t.row_count == 2
    assert t.column_count == 2
    assert t.rows[0][0].text == "Source"
    assert t.rows[1][1].text == "UN Charter Art 2(4)"

    # Verify speaker notes (OFF-001)
    notes_blocks = [b for b in doc_ir.blocks if "[Speaker Notes]" in (b.text or "")]
    assert len(notes_blocks) >= 2
    assert any("Welcome students" in b.text for b in notes_blocks)
    assert any("hard law and soft law" in b.text for b in notes_blocks)

    # Verify Arabic RTL text (LANG-001)
    ar_blocks = [b for b in doc_ir.blocks if b.language == "ar"]
    assert len(ar_blocks) >= 1
    assert ar_blocks[0].text_direction == TextDirection.RTL
    assert "مصادر القانون الدولي" in ar_blocks[0].text


def test_pptx_importer_invalid_file(tmp_path: Path):
    """Verify corrupted PPTX raises descriptive ValueError."""
    bad_file = tmp_path / "corrupt.pptx"
    bad_file.write_bytes(b"PK\x03\x04corrupted content")

    importer = PptxImporter()
    with JobWorkspace() as ws:
        with pytest.raises(ValueError, match="Failed to parse PPTX document"):
            importer.import_document(bad_file, ws)
