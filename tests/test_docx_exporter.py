"""Tests for large-print DOCX exporter (OUT-001, OUT-005..009, FN-002)."""

from pathlib import Path
import docx
from docx.shared import Mm, Pt
import pytest

from openlargeprint.exporters import (
    DocxExporter,
    ExportOptions,
    PaperSize,
    PresetName,
)
from openlargeprint.ir.models import (
    Block,
    BlockType,
    BoundingBox,
    DocumentIR,
    DocumentMetadata,
    ExtractionMethod,
    PageClassification,
    PageMetadata,
)


def create_sample_ir() -> DocumentIR:
    """Create a sample DocumentIR structure for testing export."""
    metadata = DocumentMetadata(
        title="Constitutional Law Principles",
        source_file_name="conlaw.pdf",
        page_count=2,
    )
    pages = [
        PageMetadata(page_number=1, width=595, height=842, classification=PageClassification.NATIVE),
        PageMetadata(page_number=2, width=595, height=842, classification=PageClassification.NATIVE),
    ]
    blocks = [
        Block(
            id="p1_marker",
            type=BlockType.PAGE_MARKER,
            source_page=1,
            page_marker=1,
            extraction_method=ExtractionMethod.NATIVE,
        ),
        Block(
            id="p1_b1",
            type=BlockType.HEADING,
            text="Article I: The Legislative Branch",
            level=1,
            source_page=1,
            extraction_method=ExtractionMethod.NATIVE,
        ),
        Block(
            id="p1_b2",
            type=BlockType.PARAGRAPH,
            text="All legislative Powers herein granted shall be vested in a Congress.",
            source_page=1,
            extraction_method=ExtractionMethod.NATIVE,
        ),
        Block(
            id="p1_b3",
            type=BlockType.LIST,
            text="• Qualifications of Representatives",
            source_page=1,
            extraction_method=ExtractionMethod.NATIVE,
        ),
        Block(
            id="p1_b4",
            type=BlockType.FOOTNOTE,
            text="1. Compare with early parliamentary sovereignty precedents.",
            source_page=1,
            extraction_method=ExtractionMethod.NATIVE,
        ),
        Block(
            id="p2_marker",
            type=BlockType.PAGE_MARKER,
            source_page=2,
            page_marker=2,
            extraction_method=ExtractionMethod.NATIVE,
        ),
        Block(
            id="p2_b1",
            type=BlockType.PARAGRAPH,
            text="The Senate shall be composed of two Senators from each State.",
            source_page=2,
            extraction_method=ExtractionMethod.NATIVE,
        ),
    ]
    return DocumentIR(metadata=metadata, pages=pages, blocks=blocks)


def test_docx_export_large_preset_defaults(tmp_path: Path):
    """Verify that Large preset applies 20pt body font, 1.5 line spacing, and A4 page size."""
    doc_ir = create_sample_ir()
    out_file = tmp_path / "output_large.docx"
    
    exporter = DocxExporter()
    options = ExportOptions(preset=PresetName.LARGE, paper_size=PaperSize.A4)
    exporter.export(doc_ir, out_file, options)

    assert out_file.exists()
    doc = docx.Document(str(out_file))

    # 1. Section dimensions (A4: 210mm x 297mm)
    section = doc.sections[0]
    assert pytest.approx(section.page_width.mm, 1.0) == 210.0
    assert pytest.approx(section.page_height.mm, 1.0) == 297.0

    # 2. Print reminder in metadata (OUT-009)
    assert "Print at 100%" in doc.core_properties.comments

    # 3. Check paragraphs and styles
    paras = doc.paragraphs
    # Find page marker
    marker_para = next(p for p in paras if "Original Page 1" in p.text)
    assert marker_para is not None

    # Find heading
    heading_para = next(p for p in paras if "Article I" in p.text)
    assert heading_para.runs[0].font.bold is True
    # H1 size for 20pt body is 28pt (max(26.0, 20 * 1.4))
    assert heading_para.runs[0].font.size == Pt(28)

    # Find body paragraph
    body_para = next(p for p in paras if "All legislative Powers" in p.text)
    assert body_para.runs[0].font.size == Pt(20)
    assert body_para.paragraph_format.line_spacing == 1.5

    # Find footnote (FN-002: minimum readable size enforced)
    fn_para = next(p for p in paras if "parliamentary sovereignty" in p.text)
    assert fn_para.runs[0].font.size >= Pt(14)


def test_docx_export_a3_dimensions(tmp_path: Path):
    """Verify that A3 export applies 297mm x 420mm physical dimensions (OUT-007, OUT-008)."""
    doc_ir = create_sample_ir()
    out_file = tmp_path / "output_a3.docx"

    exporter = DocxExporter()
    options = ExportOptions(preset=PresetName.EXTRA_LARGE, paper_size=PaperSize.A3)
    exporter.export(doc_ir, out_file, options)

    doc = docx.Document(str(out_file))
    section = doc.sections[0]
    assert pytest.approx(section.page_width.mm, 1.0) == 297.0
    assert pytest.approx(section.page_height.mm, 1.0) == 420.0

    # Body font for Extra Large is 24pt
    body_para = next(p for p in doc.paragraphs if "All legislative Powers" in p.text)
    assert body_para.runs[0].font.size == Pt(24)


def test_docx_export_omit_page_markers(tmp_path: Path):
    """Verify that page markers can be omitted when configured (OUT-005)."""
    doc_ir = create_sample_ir()
    out_file = tmp_path / "output_no_markers.docx"

    exporter = DocxExporter()
    options = ExportOptions(include_page_markers=False)
    exporter.export(doc_ir, out_file, options)

    doc = docx.Document(str(out_file))
    marker_paras = [p for p in doc.paragraphs if "Original Page" in p.text]
    assert len(marker_paras) == 0
