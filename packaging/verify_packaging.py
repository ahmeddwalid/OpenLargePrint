#!/usr/bin/env python3
"""Validates Tauri configuration, sidecar integration, and packaging requirements (PKG-001, PKG-002, SEC-005)."""

import json
from pathlib import Path
import sys
import subprocess
import platform
import tempfile
from build_sidecar import get_target_triple

PACKAGED_CONVERSION_TIMEOUT_SECONDS = 180


def verify_conversion(sidecar_bin: Path) -> bool:
    import docx
    import pypdfium2 as pdfium
    from PIL import Image, ImageDraw
    from reportlab.lib.pagesizes import A3, A4
    from reportlab.pdfgen import canvas

    with tempfile.TemporaryDirectory(prefix="olp-packaged-") as directory:
        root = Path(directory)
        source = root / "Three pages with spaces.pdf"
        scan = Image.new("RGB", (900, 350), "white")
        ImageDraw.Draw(scan).text((40, 80), "Judicial Review in Administrative Law", fill="black", font_size=36)
        scan_path = root / "scan.png"
        scan.save(scan_path)
        pdf = canvas.Canvas(str(source), pagesize=A4)
        for index in range(1, 4):
            if index == 2:
                pdf.drawImage(str(scan_path), 40, 300, width=500, height=195)
            else:
                pdf.drawString(72, 700, f"Native acceptance page {index} preserves exact text.")
            pdf.showPage()
        pdf.save()
        for paper, dimensions in (("A4", A4), ("A3", A3)):
            for output_format in ("pdf", "docx"):
                output = root / f"Output {paper}.{output_format}"
                command = {
                    "command": "convert", "id": f"smoke-{paper}-{output_format}",
                    "file_path": str(source), "output_path": str(output),
                    "export_format": output_format, "paper_size": paper,
                    "page_range": "2-3" if paper == "A3" else None,
                }
                run = subprocess.run(
                    [str(sidecar_bin), "sidecar"], input=json.dumps(command) + "\n",
                    capture_output=True, text=True, encoding="utf-8",
                    timeout=PACKAGED_CONVERSION_TIMEOUT_SECONDS, check=True,
                )
                events = [json.loads(line) for line in run.stdout.splitlines() if line.startswith("{")]
                success = next((event for event in events if event.get("type") == "success"), None)
                if success is None or not output.is_file():
                    return False
                if output_format == "pdf":
                    with pdfium.PdfDocument(output) as document:
                        if not len(document):
                            return False
                        texts = []
                        for index in range(len(document)):
                            page = document[index]
                            try:
                                if any(abs(actual - expected) > 1 for actual, expected in zip(page.get_size(), dimensions)):
                                    return False
                                textpage = page.get_textpage()
                                try:
                                    texts.append(textpage.get_text_range())
                                finally:
                                    textpage.close()
                                bitmap = page.render(scale=0.5)
                                bitmap.close()
                            finally:
                                page.close()
                        text = " ".join(texts)
                else:
                    document = docx.Document(output)
                    section = document.sections[0]
                    if any(abs(actual - expected) > 1 for actual, expected in zip(
                        (section.page_width.pt, section.page_height.pt), dimensions
                    )):
                        return False
                    text = " ".join(paragraph.text for paragraph in document.paragraphs)
                if "Judicial Review" not in text or "Native acceptance page 3" not in text:
                    return False
                if ("Native acceptance page 1" in text) != (paper == "A4"):
                    return False
    return True


def verify_packaging() -> bool:
    repo_root = Path(__file__).resolve().parent.parent
    tauri_conf_path = repo_root / "src-tauri" / "tauri.conf.json"

    if not tauri_conf_path.exists():
        print(f"Error: {tauri_conf_path} does not exist!")
        return False

    with open(tauri_conf_path, "r") as f:
        conf = json.load(f)

    # 1. Check frontendDist
    dist = conf.get("build", {}).get("frontendDist")
    if dist != "../ui/dist":
        print(f"Error: expected frontendDist to be '../ui/dist', got '{dist}'")
        return False

    # 2. Check CSP
    csp = conf.get("app", {}).get("security", {}).get("csp", "")
    if "default-src 'self'" not in csp:
        print(f"Error: CSP missing 'default-src self': {csp}")
        return False

    # 3. Check product details
    product_name = conf.get("productName")
    if product_name != "OpenLargePrint":
        print(f"Error: unexpected product name: {product_name}")
        return False

    # 4. Check Windows bundle targets
    targets = conf.get("bundle", {}).get("targets", [])
    if "nsis" not in targets:
        print(f"Error: bundle targets must include 'nsis' for Windows 11 packaging. Got: {targets}")
        return False

    # 5. Check externalBin declaration (PKG-001, PKG-002)
    ext_bin = conf.get("bundle", {}).get("externalBin", [])
    if "binaries/openlargeprint-sidecar" not in ext_bin:
        print(f"Error: bundle.externalBin must contain 'binaries/openlargeprint-sidecar'. Got: {ext_bin}")
        return False

    # 6. Check compiled sidecar binary exists
    suffix = ".exe" if platform.system() == "Windows" else ""
    sidecar_bin = repo_root / "src-tauri" / "binaries" / f"openlargeprint-sidecar-{get_target_triple()}{suffix}"
    if not sidecar_bin.exists():
        print(f"Error: sidecar binary not found at {sidecar_bin}. Run build_sidecar.py first.")
        return False

    try:
        result = subprocess.run(
            [str(sidecar_bin), "sidecar"],
            input='{"command":"health"}\n',
            capture_output=True,
            text=True,
            encoding="utf-8",
            timeout=120,
            check=True,
        )
        events = [json.loads(line) for line in result.stdout.splitlines() if line.strip()]
        if not any(event.get("type") == "health" and event.get("status") == "ready" for event in events):
            print("Error: packaged sidecar did not report ready.")
            return False
        if not verify_conversion(sidecar_bin):
            print("Error: packaged conversion failed OCR, page rendering, text retention, selection, or paper dimensions.")
            return False
    except (OSError, subprocess.SubprocessError, ValueError) as error:
        print(f"Error: packaged sidecar smoke test failed ({type(error).__name__}).")
        return False

    # 7. Check accessible Windows icons exist
    icons_dir = repo_root / "src-tauri" / "icons"
    if not (icons_dir / "icon.ico").exists() or not (icons_dir / "icon.png").exists():
        print(f"Error: required application icons (icon.ico, icon.png) missing in {icons_dir}.")
        return False

    # 8. Check UI build assets exist
    ui_index = repo_root / "ui" / "dist" / "index.html"
    if not ui_index.exists():
        print("Error: ui/dist/index.html not built yet. Run 'npm run build' in ui/.")
        return False

    print("Packaging smoke test passed: native/scanned multi-page conversion, A4/A3 PDF and DOCX, selected pages, icons, and frontend assets.")
    return True


if __name__ == "__main__":
    success = verify_packaging()
    sys.exit(0 if success else 1)
