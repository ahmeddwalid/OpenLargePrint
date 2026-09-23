"""Large-print DOCX exporter (OUT-001, OUT-005..009)."""

from __future__ import annotations

import re
from pathlib import Path
from typing import List, Optional
import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Mm, Pt, RGBColor

from openlargeprint.ir.models import Block, BlockType, DocumentIR, TableStructure, TextDirection
from openlargeprint.layout.table import TableTier, evaluate_table_fit
from openlargeprint.security.isolation import log_safe_info
from .base import BaseExporter, ExportOptions, PaperSize

_ILLEGAL_XML_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f]")


def clean_xml_string(text: Optional[str]) -> str:
    """Strip characters that are illegal in XML 1.0 specifications (OUT-001)."""
    if not text:
        return ""
    return _ILLEGAL_XML_CHARS_RE.sub("", text)


class DocxExporter(BaseExporter):
    """Renders DocumentIR into an accessible, single-column large-print DOCX file."""

    def export(self, doc: DocumentIR, output_path: Path, options: Optional[ExportOptions] = None) -> Path:
        """Render DocumentIR into target large-print DOCX file."""
        if options is None:
            options = ExportOptions()

        log_safe_info(
            f"Exporting DocumentIR to DOCX at {options.body_pt}pt ({options.paper_size.value}) -> {output_path.name}"
        )

        document = docx.Document()
        section = document.sections[0]

        # 1. Physical page dimensions & margins (OUT-007..009)
        self._configure_page_geometry(section, options)
        settings = document.settings.element
        view = settings.find(qn("w:view"))
        if view is None:
            view = OxmlElement("w:view")
            settings.insert(0, view)
        view.set(qn("w:val"), "print")
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("Page ")
        for field_name, separator in (("PAGE", " of "), ("NUMPAGES", "")):
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), field_name)
            footer._p.append(field)
            if separator:
                footer.add_run(separator)

        # 2. Add print reminder notice (OUT-009)
        core_props = document.core_properties
        core_props.title = clean_xml_string(doc.metadata.title) or "OpenLargePrint Document"
        core_props.comments = (
            f"OpenLargePrint output: {options.body_pt}pt text, {options.paper_size.value} paper size. "
            "REMINDER: Print at 100% / actual size rather than 'fit to page' to preserve text size."
        )

        # 3. Render blocks sequentially in single-column reflow
        usable_width = section.page_width - section.left_margin - section.right_margin
        usable_height = section.page_height - section.top_margin - section.bottom_margin

        for block in doc.blocks:
            self._render_block(document, block, options, usable_width, usable_height)

        # Ensure parent directory exists
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(output_path))
        log_safe_info(f"DOCX export completed successfully: {output_path.name}")
        return output_path

    def _configure_page_geometry(self, section: docx.section.Section, options: ExportOptions) -> None:
        """Set true physical page dimensions and proportional margins (OUT-007..009)."""
        if options.paper_size == PaperSize.A3:
            section.page_width = Mm(297)
            section.page_height = Mm(420)
            section.top_margin = Mm(25)
            section.bottom_margin = Mm(25)
            section.left_margin = Mm(25)
            section.right_margin = Mm(25)
        else:
            # Default A4
            section.page_width = Mm(210)
            section.page_height = Mm(297)
            section.top_margin = Mm(20)
            section.bottom_margin = Mm(20)
            section.left_margin = Mm(20)
            section.right_margin = Mm(20)

    def _get_or_create_monochrome_image(self, file_path: str) -> str:
        """Convert image to high-contrast grayscale for laser printing."""
        try:
            from PIL import Image, ImageEnhance
            src_p = Path(file_path)
            mono_path = src_p.parent / f"{src_p.stem}_mono.png"
            if mono_path.exists():
                return str(mono_path)
            with Image.open(src_p) as im:
                gray = im.convert("L")
                enhancer = ImageEnhance.Contrast(gray)
                enhanced = enhancer.enhance(1.15)
                enhanced.save(mono_path, format="PNG")
                return str(mono_path)
        except Exception:
            return file_path

    def _render_block(
        self,
        document: docx.Document,
        block: Block,
        options: ExportOptions,
        usable_width: Inches,
        usable_height: Optional[Inches] = None,
    ) -> None:
        """Render an individual semantic block to DOCX according to typography preset."""
        safe_text = clean_xml_string(block.text)
        is_mono = options.monochrome

        # Handle page markers (OUT-005)
        if block.type == BlockType.PAGE_MARKER:
            if options.include_page_markers and block.page_marker is not None:
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(options.body_pt * 0.8)
                p.paragraph_format.space_after = Pt(options.body_pt * 0.5)
                run = p.add_run(f"— Original Page {block.page_marker} —")
                run.font.name = options.font_family
                run.font.size = Pt(max(12.0, options.body_pt * 0.7))
                run.font.bold = True
                run.font.color.rgb = RGBColor(0, 0, 0) if is_mono else RGBColor(90, 90, 90)
            return

        # Handle headings
        if block.type in (BlockType.TITLE, BlockType.HEADING):
            level = block.level or 1
            p = document.add_paragraph()
            
            if level == 1:
                size_pt = max(26.0, options.body_pt * 1.4)
                p.paragraph_format.space_before = Pt(20)
                p.paragraph_format.space_after = Pt(10)
            elif level == 2:
                size_pt = max(23.0, options.body_pt * 1.25)
                p.paragraph_format.space_before = Pt(16)
                p.paragraph_format.space_after = Pt(8)
            else:
                size_pt = max(21.0, options.body_pt * 1.15)
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(6)

            p.paragraph_format.keep_with_next = True
            run = p.add_run(safe_text)
            run.font.name = options.font_family
            run.font.size = Pt(size_pt)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0) if is_mono else RGBColor(20, 20, 20)
            self._apply_text_direction(p, run, block, options)
            return

        # Handle lists
        if block.type == BlockType.LIST:
            p = document.add_paragraph(style="List Bullet")
            p.paragraph_format.line_spacing = options.line_spacing
            p.paragraph_format.space_after = Pt(options.body_pt * 0.3)
            # Strip initial bullet if text already has one
            raw_text = safe_text.lstrip("•-* \t")
            run = p.add_run(raw_text)
            run.font.name = options.font_family
            run.font.size = Pt(options.body_pt)
            if is_mono:
                run.font.color.rgb = RGBColor(0, 0, 0)
            self._apply_text_direction(p, run, block, options)
            return

        # Handle quotes
        if block.type == BlockType.QUOTE:
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.line_spacing = options.line_spacing
            p.paragraph_format.space_after = Pt(options.body_pt * 0.5)
            run = p.add_run(safe_text)
            run.font.name = options.font_family
            run.font.size = Pt(options.body_pt)
            run.font.italic = True
            if is_mono:
                run.font.color.rgb = RGBColor(0, 0, 0)
            self._apply_text_direction(p, run, block, options)
            return

        # Handle captions
        if block.type == BlockType.CAPTION:
            min_size = max(14.0, options.body_pt * 0.85)
            p = document.add_paragraph()
            p.alignment = (
                WD_ALIGN_PARAGRAPH.RIGHT
                if block.text_direction == TextDirection.RTL
                else WD_ALIGN_PARAGRAPH.CENTER
            )
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(10)
            run = p.add_run(safe_text)
            run.font.name = options.font_family
            run.font.size = Pt(min_size)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0) if is_mono else RGBColor(60, 60, 60)
            self._apply_text_direction(p, run, block, options)
            return

        # Handle footnotes (FN-001, FN-002: enforce minimum readable size >= 14pt)
        if block.type == BlockType.FOOTNOTE:
            min_size = max(14.0, options.body_pt * 0.80)
            p = document.add_paragraph()
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(safe_text)
            run.font.name = options.font_family
            run.font.size = Pt(min_size)
            run.font.italic = True
            run.font.color.rgb = RGBColor(0, 0, 0) if is_mono else RGBColor(70, 70, 70)
            self._apply_text_direction(p, run, block, options)
            return

        # Handle tables (TBL-001, TBL-002, FN-002)
        if block.type == BlockType.TABLE and block.table_structure:
            self._render_table(document, block, options, usable_width, usable_height)
            return

        # Handle images (IMG-001, IMG-003)
        if block.type == BlockType.IMAGE and block.image_asset:
            asset = block.image_asset
            if asset.file_path and Path(asset.file_path).exists():
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(12)

                # Fit within usable width AND height without distortion (IMG-003)
                native_w_emu = int(asset.width / 96.0 * 914400)
                native_h_emu = int(asset.height / 96.0 * 914400)
                disp_w = min(int(usable_width), native_w_emu)
                disp_h = int(disp_w * (native_h_emu / max(1, native_w_emu)))
                if usable_height is not None and int(usable_height) > 0 and disp_h > int(usable_height):
                    disp_h = int(usable_height)
                    disp_w = int(disp_h * (native_w_emu / max(1, native_h_emu)))

                img_path = asset.file_path
                if options.monochrome:
                    img_path = self._get_or_create_monochrome_image(asset.file_path)

                p.add_run().add_picture(img_path, width=Emu(disp_w), height=Emu(disp_h))
            return

        # Default: Regular body paragraph
        p = document.add_paragraph()
        p.paragraph_format.line_spacing = options.line_spacing
        p.paragraph_format.space_after = Pt(options.body_pt * 0.55)  # Generous paragraph spacing
        run = p.add_run(safe_text)
        run.font.name = options.font_family
        run.font.size = Pt(options.body_pt)
        run.font.color.rgb = RGBColor(0, 0, 0) if is_mono else RGBColor(20, 20, 20)
        self._apply_text_direction(p, run, block, options)

    def _apply_text_direction(self, p, run, block: Block, options: ExportOptions) -> None:
        """Apply RTL paragraph alignment and OpenXML attributes if block is RTL (LANG-001)."""
        if block.text_direction == TextDirection.RTL:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            pPr = p._p.get_or_add_pPr()
            if pPr.find(qn("w:bidi")) is None:
                pPr.append(OxmlElement("w:bidi"))

            rPr = run._r.get_or_add_rPr()
            if rPr.find(qn("w:rtl")) is None:
                rPr.append(OxmlElement("w:rtl"))

            # Complex script font for Arabic
            rFonts = rPr.find(qn("w:rFonts"))
            if rFonts is None:
                rFonts = OxmlElement("w:rFonts")
                rPr.append(rFonts)
            rFonts.set(qn("w:cs"), "Traditional Arabic")
            rFonts.set(qn("w:ascii"), options.font_family)
            rFonts.set(qn("w:hAnsi"), options.font_family)

    def _render_table(
        self,
        document: docx.Document,
        block: Block,
        options: ExportOptions,
        usable_width,
        usable_height=None,
    ) -> None:
        """Render table with large-print scaling, fallback cascade, and RTL support (TBL-001, TBL-002, FN-002)."""
        table_struct = block.table_structure
        if not table_struct or not table_struct.rows:
            return

        usable_w_pt = usable_width.pt if hasattr(usable_width, "pt") else float(usable_width) / 12700.0
        min_readable_size = max(14.0, options.body_pt * 0.8)
        fit = evaluate_table_fit(
            table_struct,
            available_width=usable_w_pt,
            font_pt=options.body_pt,
            min_readable_pt=min_readable_size,
        )

        # Render visible warning if table required fallback/linearization (TBL-002)
        if fit.warning or block.warnings:
            all_warnings = list(block.warnings)
            if fit.warning and fit.warning not in all_warnings:
                all_warnings.append(fit.warning)
            for warn in all_warnings:
                p_warn = document.add_paragraph()
                p_warn.paragraph_format.space_before = Pt(8)
                p_warn.paragraph_format.space_after = Pt(4)
                r_warn = p_warn.add_run(f"⚠️ [Table Note] {clean_xml_string(warn)}")
                r_warn.font.name = options.font_family
                r_warn.font.size = Pt(max(13.0, min_readable_size * 0.9))
                r_warn.font.italic = True
                r_warn.font.color.rgb = (
                    RGBColor(0, 0, 0) if options.monochrome else RGBColor(160, 80, 0)
                )
                self._apply_text_direction(p_warn, r_warn, block, options)

        # Tier 3: Linearized accessible cards (TBL-001)
        if fit.tier == TableTier.LINEARIZE:
            # Retain source table crop image if available (TBL-001)
            if block.image_asset and block.image_asset.file_path and Path(block.image_asset.file_path).exists():
                p_img = document.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(8)
                retained_path = block.image_asset.file_path
                if options.monochrome:
                    retained_path = self._get_or_create_monochrome_image(retained_path)
                native_w_emu = int(block.image_asset.width / 96.0 * 914400)
                native_h_emu = int(block.image_asset.height / 96.0 * 914400)
                disp_w = min(int(usable_width), native_w_emu)
                disp_h = int(disp_w * (native_h_emu / max(1, native_w_emu)))
                if usable_height is not None and int(usable_height) > 0 and disp_h > int(usable_height):
                    disp_h = int(usable_height)
                    disp_w = int(disp_h * (native_w_emu / max(1, native_h_emu)))
                p_img.add_run().add_picture(retained_path, width=Emu(disp_w), height=Emu(disp_h))

            lines = (fit.linearized_text or table_struct.to_linearized_text()).split("\n")
            for line in lines:
                if not line.strip():
                    continue
                p = document.add_paragraph()
                p.paragraph_format.line_spacing = 1.3
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(clean_xml_string(line))
                run.font.name = options.font_family
                run.font.size = Pt(min_readable_size)
                if line.startswith(("•", "[Table")):
                    run.font.bold = True
                self._apply_text_direction(p, run, block, options)
            return

        # Tier 2: Split columns sub-tables (TBL-001)
        if fit.tier == TableTier.SPLIT and fit.split_tables:
            for st in fit.split_tables:
                self._build_docx_table(document, st, options, usable_width, block.text_direction)
                document.add_paragraph().paragraph_format.space_after = Pt(10)
            return

        # Tier 1: Enlarged semantic table
        self._build_docx_table(
            document,
            table_struct,
            options,
            usable_width,
            block.text_direction,
            fit.column_widths,
        )
        document.add_paragraph().paragraph_format.space_after = Pt(10)

    def _build_docx_table(
        self,
        document: docx.Document,
        table_struct: TableStructure,
        options: ExportOptions,
        usable_width,
        text_direction: TextDirection,
        col_widths_pt: Optional[List[float]] = None,
    ) -> None:
        """Construct styled python-docx table with headers, shading, cantSplit, and bidiVisual."""
        if not table_struct.rows:
            return

        cols = table_struct.column_count
        tbl = document.add_table(rows=len(table_struct.rows), cols=cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Solid black borders for monochrome laser printing (DESIGN.md §8.1)
        if options.monochrome:
            tblPr = tbl._tbl.tblPr
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                el = OxmlElement(f"w:{edge}")
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "8")
                el.set(qn("w:space"), "0")
                el.set(qn("w:color"), "000000")
                borders.append(el)
            tblPr.append(borders)

        # RTL visual ordering (LANG-001)
        if text_direction == TextDirection.RTL:
            tblPr = tbl._tbl.tblPr
            if tblPr.find(qn("w:bidiVisual")) is None:
                tblPr.append(OxmlElement("w:bidiVisual"))

        cell_font_size = Pt(max(14.0, options.body_pt * 0.85))

        for r_idx, row in enumerate(table_struct.rows):
            tr = tbl.rows[r_idx]
            trPr = tr._tr.get_or_add_trPr()

            # Repeat header across pages
            if r_idx == 0 and table_struct.has_header:
                if trPr.find(qn("w:tblHeader")) is None:
                    trPr.append(OxmlElement("w:tblHeader"))

            # Prevent the header row from splitting, but let body rows break across
            # pages so tall rows reflow instead of clipping (TBL-001, large-print).
            is_header_row = r_idx == 0 and table_struct.has_header
            if is_header_row and trPr.find(qn("w:cantSplit")) is None:
                trPr.append(OxmlElement("w:cantSplit"))

            for c_idx, cell_data in enumerate(row):
                if c_idx >= cols:
                    break
                cell = tbl.cell(r_idx, c_idx)
                # Set column width if calculated
                if col_widths_pt and c_idx < len(col_widths_pt):
                    cell.width = Pt(col_widths_pt[c_idx])

                # Header shading: white in monochrome (black text stays readable),
                # light gray otherwise (DESIGN.md §8.1).
                if r_idx == 0 and table_struct.has_header:
                    tcPr = cell._tc.get_or_add_tcPr()
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), "FFFFFF" if options.monochrome else "F2F2F2")
                    tcPr.append(shd)

                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.2
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                run = p.add_run(clean_xml_string(cell_data.text))
                run.font.name = options.font_family
                run.font.size = cell_font_size
                if r_idx == 0 and table_struct.has_header:
                    run.font.bold = True

                # Apply text direction inside cell
                dummy_block = Block(
                    id="cell",
                    type=BlockType.PARAGRAPH,
                    source_page=1,
                    text_direction=text_direction,
                )
                self._apply_text_direction(p, run, dummy_block, options)
