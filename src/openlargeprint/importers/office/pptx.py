"""Native PPTX importer normalizing slides, shapes, tables, and speaker notes into DocumentIR (OFF-001, DOC-001)."""

from __future__ import annotations

import re
from pathlib import Path
from typing import List, Optional

import pptx
from PIL import Image

from openlargeprint.importers.base import BaseImporter
from openlargeprint.ir.models import (
    Block,
    BlockType,
    DocumentIR,
    DocumentMetadata,
    ExtractionMethod,
    ImageAsset,
    PageClassification,
    PageMetadata,
    TableCell,
    TableStructure,
    TextDirection,
)
from openlargeprint.security import JobWorkspace, validate_image_dimensions
from openlargeprint.text.direction import detect_language, detect_text_direction


from openlargeprint.importers.base import BaseImporter, CancelCheck, CheckpointCallback, ProgressCallback


class PptxImporter(BaseImporter):
    """Imports PowerPoint (.pptx) presentations into canonical DocumentIR."""

    def import_document(
        self,
        file_path: Path,
        workspace: JobWorkspace,
        progress_callback: Optional[ProgressCallback] = None,
        cancel_check: Optional[CancelCheck] = None,
        checkpoint_callback: Optional[CheckpointCallback] = None,
    ) -> DocumentIR:
        """Parse PowerPoint (.pptx) document and return normalized DocumentIR."""
        path = Path(file_path).resolve()
        try:
            prs = pptx.Presentation(path)
        except Exception as e:
            raise ValueError(f"Failed to parse PPTX document: {e}") from e

        blocks: List[Block] = []
        pages: List[PageMetadata] = []
        img_counter = 0
        document_warnings: List[str] = []

        # Slide dimensions in points (72 points per inch; pptx units are EMUs: 1 pt = 12700 EMUs)
        slide_w_pt = round(float(prs.slide_width) / 12700.0, 2)
        slide_h_pt = round(float(prs.slide_height) / 12700.0, 2)

        total_slides = len(prs.slides)
        first_title: Optional[str] = None

        for slide_idx, slide in enumerate(prs.slides, start=1):
            if cancel_check and cancel_check():
                raise InterruptedError("Operation cancelled by user.")

            if progress_callback:
                progress_callback(
                    slide_idx,
                    total_slides,
                    "extracting",
                    f"Extracting slide presentation — slide {slide_idx} of {total_slides}",
                )

            # 1. Slide PageMetadata
            pages.append(
                PageMetadata(
                    page_number=slide_idx,
                    width=slide_w_pt,
                    height=slide_h_pt,
                    classification=PageClassification.NATIVE,
                )
            )

            # 2. Slide Transition Marker (OUT-005)
            blocks.append(
                Block(
                    id=f"slide_{slide_idx}_marker",
                    type=BlockType.PAGE_MARKER,
                    page_marker=slide_idx,
                    source_page=slide_idx,
                    extraction_method=ExtractionMethod.NATIVE,
                )
            )

            # 3. Slide Title (Heading 1)
            title_shape = slide.shapes.title
            if title_shape and title_shape.has_text_frame:
                title_text = title_shape.text_frame.text.strip()
                if title_text:
                    if first_title is None:
                        first_title = title_text
                    dir_t = detect_text_direction(title_text)
                    lang_t = detect_language(title_text)
                    blocks.append(
                        Block(
                            id=f"slide_{slide_idx}_title",
                            type=BlockType.HEADING,
                            text=title_text,
                            level=1,
                            language=lang_t,
                            text_direction=dir_t,
                            source_page=slide_idx,
                            extraction_method=ExtractionMethod.NATIVE,
                        )
                    )

            # 4. Sort shapes top-to-bottom for natural single-column reading order
            sorted_shapes = []
            for shape in slide.shapes:
                if shape == title_shape:
                    continue
                top = getattr(shape, "top", 0) or 0
                left = getattr(shape, "left", 0) or 0
                sorted_shapes.append((top, left, shape))
            sorted_shapes.sort(key=lambda s: (s[0], s[1]))

            # 5. Extract shapes
            for s_idx, (_, _, shape) in enumerate(sorted_shapes):
                # Images
                if hasattr(shape, "image"):
                    try:
                        img_counter += 1
                        img_blob = shape.image.blob
                        ext = getattr(shape.image, "ext", "png")
                        img_filename = f"pptx_img_{img_counter}.{ext}"
                        img_path = workspace.assets_dir / img_filename
                        img_path.write_bytes(img_blob)

                        with Image.open(img_path) as pil_img:
                            w, h = pil_img.size
                            validate_image_dimensions(w, h)

                        asset = ImageAsset(
                            asset_id=f"pptx_asset_{img_counter}",
                            file_path=str(img_path),
                            width=w,
                            height=h,
                            mime_type=f"image/{ext}",
                        )
                        blocks.append(
                            Block(
                                id=f"slide_{slide_idx}_img_{s_idx}",
                                type=BlockType.IMAGE,
                                image_asset=asset,
                                source_page=slide_idx,
                                extraction_method=ExtractionMethod.OFFICE_IMPORT,
                            )
                        )
                    except Exception as exc:
                        # Never drop a figure silently.
                        document_warnings.append(
                            f"A slide image could not be imported ({type(exc).__name__}); "
                            "it may be missing from the output"
                        )
                    continue

                # Tables (TBL-001)
                if shape.has_table:
                    tbl_block = self._process_pptx_table(shape.table, slide_idx, s_idx)
                    if tbl_block:
                        blocks.append(tbl_block)
                    continue

                # Text frames
                if shape.has_text_frame:
                    tf_blocks = self._process_text_frame(shape.text_frame, slide_idx, s_idx)
                    blocks.extend(tf_blocks)

            # 6. Speaker Notes (OFF-001)
            if slide.has_notes_slide:
                try:
                    notes_frame = slide.notes_slide.notes_text_frame
                    if notes_frame:
                        notes_text = notes_frame.text.strip()
                        # Filter out default placeholder / slide number strings
                        if notes_text and not notes_text.isdigit():
                            dir_n = detect_text_direction(notes_text)
                            lang_n = detect_language(notes_text)
                            blocks.append(
                                Block(
                                    id=f"slide_{slide_idx}_notes",
                                    type=BlockType.PARAGRAPH,
                                    text=f"[Speaker Notes] {notes_text}",
                                    language=lang_n,
                                    text_direction=dir_n,
                                    source_page=slide_idx,
                                    extraction_method=ExtractionMethod.NATIVE,
                                )
                            )
                except Exception as exc:
                    # Never drop speaker notes silently (OFF-001, SPEC §2.3).
                    document_warnings.append(
                        f"Speaker notes on slide {slide_idx} could not be imported "
                        f"({type(exc).__name__}); they may be missing from the output"
                    )

        # Surface anything that could not be imported (never drop silently).
        if document_warnings:
            target = next((b for b in blocks if b.type != BlockType.PAGE_MARKER), None)
            if target is not None:
                target.warnings.extend(document_warnings)

        doc_title = first_title or path.stem.replace("_", " ").title()

        metadata = DocumentMetadata(
            title=doc_title,
            page_count=max(1, total_slides),
            source_file_name=path.name,
        )

        return DocumentIR(
            schema_version="1.0.0",
            metadata=metadata,
            pages=pages,
            blocks=blocks,
        )

    def _process_text_frame(
        self, text_frame, slide_idx: int, shape_idx: int
    ) -> List[Block]:
        """Convert paragraphs in a PowerPoint text frame into BlockType.LIST or PARAGRAPH."""
        blocks: List[Block] = []

        for p_idx, p in enumerate(text_frame.paragraphs):
            raw_text = p.text.strip()
            if not raw_text:
                continue

            direction = detect_text_direction(raw_text)
            lang = detect_language(raw_text)
            block_id = f"slide_{slide_idx}_s{shape_idx}_p{p_idx}"

            # Check if paragraph is bulleted / indented list item
            level = getattr(p, "level", 0) or 0
            is_bullet = level > 0 or raw_text.startswith(("•", "-", "*"))

            if is_bullet:
                clean_text = raw_text.lstrip("•-* \t")
                blocks.append(
                    Block(
                        id=block_id,
                        type=BlockType.LIST,
                        text=clean_text if clean_text else raw_text,
                        language=lang,
                        text_direction=direction,
                        source_page=slide_idx,
                        extraction_method=ExtractionMethod.NATIVE,
                    )
                )
            else:
                blocks.append(
                    Block(
                        id=block_id,
                        type=BlockType.PARAGRAPH,
                        text=raw_text,
                        language=lang,
                        text_direction=direction,
                        source_page=slide_idx,
                        extraction_method=ExtractionMethod.NATIVE,
                    )
                )

        return blocks

    def _process_pptx_table(
        self, table, slide_idx: int, shape_idx: int
    ) -> Optional[Block]:
        """Convert python-pptx Table into canonical TableStructure (TBL-001)."""
        if not table.rows:
            return None

        table_rows: List[List[TableCell]] = []
        is_rtl_table = False

        for r_idx, row in enumerate(table.rows):
            cells_row: List[TableCell] = []
            is_header = (r_idx == 0)

            for cell in row.cells:
                text = cell.text.strip()
                cell_dir = detect_text_direction(text)
                if cell_dir == TextDirection.RTL:
                    is_rtl_table = True

                cells_row.append(
                    TableCell(
                        text=text,
                        is_header=is_header,
                    )
                )
            table_rows.append(cells_row)

        table_struct = TableStructure(
            has_header=True,
            rows=table_rows,
        )

        direction = TextDirection.RTL if is_rtl_table else TextDirection.LTR
        lang = "ar" if is_rtl_table else "en"

        return Block(
            id=f"slide_{slide_idx}_tbl_{shape_idx}",
            type=BlockType.TABLE,
            text=table_struct.to_markdown_table(),
            table_structure=table_struct,
            language=lang,
            text_direction=direction,
            source_page=slide_idx,
            extraction_method=ExtractionMethod.NATIVE,
        )
